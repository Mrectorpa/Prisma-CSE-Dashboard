"""
Excel workbook generation for the certificate expiry report.

Produces one workbook containing:
  - A "Summary" tab: one row per tenant with counts of Red/Yellow/Green/
    Expired certificates.
  - One tab per tenant with a reportable certificate, listing each
    certificate's Name, Common Name, Subject, Issuer, Not Valid After,
    Days Until Expiry, Status, CA flag, and Folder.

Tenant tab names are sanitized to satisfy Excel's sheet-name rules
(<=31 chars, no \\ / ? * [ ] characters) and de-duplicated by appending
the tenant id when two tenants would otherwise produce the same name.
"""
import io
import re
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .classifier import ClassifiedCertificate
from .config import STATUS_EXPIRED, STATUS_GREEN, STATUS_RED, STATUS_YELLOW
from .orchestrator import OrchestrationResult, TenantReportData

# Characters Excel forbids in sheet names.
_INVALID_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")
_MAX_SHEET_NAME_LEN = 31

_TENANT_COLUMNS = [
    "Certificate Name",
    "Common Name",
    "Subject",
    "Issuer",
    "Not Valid After",
    "Days Until Expiry",
    "Status",
    "CA",
    "Folder",
]

_SUMMARY_COLUMNS = [
    "Tenant",
    "Tenant ID",
    "Expired",
    "Red (<=30 days)",
    "Yellow (31-60 days)",
    "Green (61-90 days)",
    "Total",
]

_HEADER_FONT = Font(bold=True)

# Row order within each tenant tab: most urgent first.
_STATUS_SORT_ORDER = {
    STATUS_EXPIRED: 0,
    STATUS_RED: 1,
    STATUS_YELLOW: 2,
    STATUS_GREEN: 3,
}

# Solid background fill colors per status, applied to entire data rows
# in each tenant tab. Expired uses a distinct dark red/maroon so it
# reads as more urgent than the <=30-day Red bucket.
_STATUS_FILLS = {
    STATUS_EXPIRED: PatternFill(start_color="4A0000", end_color="4A0000", fill_type="solid"),
    STATUS_RED: PatternFill(start_color="F8CBCB", end_color="F8CBCB", fill_type="solid"),
    STATUS_YELLOW: PatternFill(start_color="FDECC8", end_color="FDECC8", fill_type="solid"),
    STATUS_GREEN: PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"),
}

# Font color to use on top of each fill (white text on the dark
# Expired background for readability; default/black elsewhere).
_STATUS_FONTS = {
    STATUS_EXPIRED: Font(color="FFFFFF"),
}


def _sanitize_sheet_name(display_name: str, tenant_id: str, used_names: Dict[str, int]) -> str:
    """
    Produce a valid, unique Excel sheet name derived from a tenant's
    display name, truncating and appending the tenant id on collision.
    """
    cleaned = _INVALID_SHEET_CHARS.sub("_", display_name).strip() or tenant_id
    cleaned = cleaned[:_MAX_SHEET_NAME_LEN]

    candidate = cleaned
    if candidate.lower() not in used_names:
        used_names[candidate.lower()] = 1
        return candidate

    # Collision: append the tenant id, truncating the base name so the
    # combined result still fits within Excel's 31-character limit.
    suffix = f"_{tenant_id}"
    truncated_base = cleaned[: max(0, _MAX_SHEET_NAME_LEN - len(suffix))]
    candidate = f"{truncated_base}{suffix}"[:_MAX_SHEET_NAME_LEN]

    # In the extremely unlikely event this ALSO collides, fall back to
    # a numeric suffix to guarantee uniqueness.
    base_candidate = candidate
    counter = 2
    while candidate.lower() in used_names:
        suffix2 = f"_{counter}"
        candidate = f"{base_candidate[: _MAX_SHEET_NAME_LEN - len(suffix2)]}{suffix2}"
        counter += 1

    used_names[candidate.lower()] = 1
    return candidate


def _write_header_row(ws: Worksheet, columns: List[str]) -> None:
    for col_idx, column_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column_name)
        cell.font = _HEADER_FONT


def _autosize_columns(ws: Worksheet, columns: List[str], min_width: int = 12, max_width: int = 60) -> None:
    """Rough column auto-sizing based on the longest value seen per column."""
    widths = [len(name) for name in columns]
    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row):
            value_len = len(str(cell.value)) if cell.value is not None else 0
            if value_len > widths[idx]:
                widths[idx] = value_len

    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(
            max(width + 2, min_width), max_width
        )


def _write_summary_sheet(wb: Workbook, tenant_reports: List[TenantReportData]) -> None:
    ws: Worksheet = wb.active  # type: ignore[assignment]  # a fresh Workbook always has an active sheet
    ws.title = "Summary"
    _write_header_row(ws, _SUMMARY_COLUMNS)

    row_idx = 2
    for report in tenant_reports:
        counts = {
            STATUS_EXPIRED: 0,
            STATUS_RED: 0,
            STATUS_YELLOW: 0,
            STATUS_GREEN: 0,
        }
        for classified in report.classified_certificates:
            counts[classified.status] += 1

        total = sum(counts.values())

        ws.cell(row=row_idx, column=1, value=report.tenant.display_name)
        ws.cell(row=row_idx, column=2, value=report.tenant.id)
        ws.cell(row=row_idx, column=3, value=counts[STATUS_EXPIRED])
        ws.cell(row=row_idx, column=4, value=counts[STATUS_RED])
        ws.cell(row=row_idx, column=5, value=counts[STATUS_YELLOW])
        ws.cell(row=row_idx, column=6, value=counts[STATUS_GREEN])
        ws.cell(row=row_idx, column=7, value=total)
        row_idx += 1

    _autosize_columns(ws, _SUMMARY_COLUMNS)


def _write_tenant_sheet(wb: Workbook, report: TenantReportData, sheet_name: str) -> None:
    ws = wb.create_sheet(title=sheet_name)
    _write_header_row(ws, _TENANT_COLUMNS)

    # Most urgent certificates first (Expired/Red before Yellow/Green),
    # then by soonest expiry within the same status.
    sorted_certs: List[ClassifiedCertificate] = sorted(
        report.classified_certificates,
        key=lambda c: (
            _STATUS_SORT_ORDER.get(c.status, 99),
            c.certificate.days_until_expiry
            if c.certificate.days_until_expiry is not None
            else float("inf"),
        ),
    )

    row_idx = 2
    for classified in sorted_certs:
        cert = classified.certificate
        not_valid_after_display = (
            cert.not_valid_after.strftime("%Y-%m-%d") if cert.not_valid_after else cert.not_valid_after_raw
        )

        row_values: List[object] = [
            cert.name,
            cert.common_name,
            cert.subject,
            cert.issuer,
            not_valid_after_display,
            cert.days_until_expiry,
            classified.status,
            "Yes" if cert.is_ca else "No",
            cert.folder,
        ]

        fill = _STATUS_FILLS.get(classified.status)
        font = _STATUS_FONTS.get(classified.status)

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font

        row_idx += 1

    ws.freeze_panes = "A2"
    _autosize_columns(ws, _TENANT_COLUMNS)


def build_workbook(result: OrchestrationResult) -> bytes:
    """
    Build the full .xlsx workbook (Summary + one tab per tenant) and
    return it as raw bytes, ready to be streamed as a file download.
    """
    wb = Workbook()

    _write_summary_sheet(wb, result.tenant_reports)

    used_sheet_names: Dict[str, int] = {"summary": 1}
    for report in result.tenant_reports:
        sheet_name = _sanitize_sheet_name(
            report.tenant.display_name, report.tenant.id, used_sheet_names
        )
        _write_tenant_sheet(wb, report, sheet_name)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
